"""
mapear_pastas.py
────────────────
Lê a hierarquia de pastas a partir de um diretório raiz e gera:
  1. Um arquivo TXT com a árvore visual de pastas
  2. Um arquivo Excel com a hierarquia em colunas (Nível 1, Nível 2, Nível 3...)

Como usar:
  1. Coloque este script em qualquer lugar
  2. Edite a variável PASTA_RAIZ abaixo com o caminho da sua pasta de patentes
  3. Execute: python mapear_pastas.py
  4. Os arquivos de saída serão gerados na mesma pasta do script
"""

import os
import pathlib
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# ══════════════════════════════════════════════════════════════
# CONFIGURAÇÃO — edite apenas esta linha
# ══════════════════════════════════════════════════════════════
PASTA_RAIZ = r"C:\caminho\para\sua\pasta"   # ← edite aqui
# ══════════════════════════════════════════════════════════════

# Arquivos de saída (gerados na mesma pasta do script)
SCRIPT_DIR  = pathlib.Path(__file__).parent
SAIDA_TXT   = SCRIPT_DIR / "hierarquia_pastas.txt"
SAIDA_EXCEL = SCRIPT_DIR / "hierarquia_pastas.xlsx"


# ── 1. Coleta a hierarquia de pastas ──────────────────────────
def coletar_pastas(raiz: str) -> list[dict]:
    """
    Percorre a árvore de diretórios e retorna uma lista de dicts:
      { 'caminho': Path, 'nome': str, 'nivel': int, 'partes': list[str] }
    Ignora arquivos; só pastas.
    """
    raiz_path = pathlib.Path(raiz)
    if not raiz_path.exists():
        print(f"\n  ✖ Pasta não encontrada: {raiz}")
        print("  Edite a variável PASTA_RAIZ no início do script.\n")
        raise SystemExit(1)

    resultado = []
    for dirpath, dirnames, _ in os.walk(raiz_path):
        dirnames.sort()   # ordem alfabética
        p      = pathlib.Path(dirpath)
        nivel  = len(p.relative_to(raiz_path).parts)
        partes = list(p.relative_to(raiz_path).parts)
        resultado.append({
            'caminho': p,
            'nome'   : p.name if nivel > 0 else raiz_path.name,
            'nivel'  : nivel,
            'partes' : partes,
        })

    return resultado


# ── 2. Gera arquivo TXT com árvore visual ─────────────────────
def gerar_txt(pastas: list[dict], saida: pathlib.Path, raiz_nome: str):
    linhas = [f"📁 {raiz_nome}\n"]
    for p in pastas:
        if p['nivel'] == 0:
            continue
        indent = "│   " * (p['nivel'] - 1) + "├── "
        linhas.append(f"{indent}📁 {p['nome']}\n")

    with open(saida, "w", encoding="utf-8") as f:
        f.writelines(linhas)

    print(f"  ✔ TXT gerado  : {saida}")


# ── 3. Gera Excel com hierarquia em colunas ───────────────────
def gerar_excel(pastas: list[dict], saida: pathlib.Path, raiz_nome: str):
    wb = Workbook()
    ws = wb.active
    ws.title = "Hierarquia"

    # Descobre a profundidade máxima
    max_nivel = max((p['nivel'] for p in pastas), default=0)

    # ── Cores por nível ───────────────────────────────────────
    CORES = [
        "2C3E50",  # nível 0 (raiz)  — cinza escuro
        "1A5276",  # nível 1         — azul escuro
        "1F618D",  # nível 2         — azul médio
        "2874A6",  # nível 3         — azul claro
        "3498DB",  # nível 4         — azul mais claro
        "5DADE2",  # nível 5+        — azul suave
    ]

    def cor(nivel: int) -> str:
        return CORES[min(nivel, len(CORES) - 1)]

    def thin():
        s = Side(style='thin', color="CCCCCC")
        return Border(left=s, right=s, top=s, bottom=s)

    # ── Cabeçalho ─────────────────────────────────────────────
    headers = ["Nível " + str(i) for i in range(max_nivel + 1)]
    headers[0] = "Raiz"
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.font      = Font(bold=True, color="FFFFFF", size=11)
        c.fill      = PatternFill("solid", fgColor="1C1C1C")
        c.alignment = Alignment(horizontal='center', vertical='center')
        c.border    = thin()
        ws.column_dimensions[c.column_letter].width = 28

    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"

    # ── Dados ─────────────────────────────────────────────────
    # Linha da raiz
    row = 2
    c = ws.cell(row=row, column=1, value=raiz_nome)
    c.font      = Font(bold=True, color="FFFFFF", size=10)
    c.fill      = PatternFill("solid", fgColor=cor(0))
    c.alignment = Alignment(horizontal='left', vertical='center',
                             indent=1)
    c.border    = thin()
    row += 1

    for p in pastas:
        if p['nivel'] == 0:
            continue
        col = p['nivel'] + 1   # nível 1 → coluna 2
        c = ws.cell(row=row, column=col, value=p['nome'])
        c.font      = Font(color="FFFFFF", size=10,
                           bold=(p['nivel'] == 1))
        c.fill      = PatternFill("solid", fgColor=cor(p['nivel']))
        c.alignment = Alignment(horizontal='left', vertical='center',
                                 indent=1)
        c.border    = thin()
        ws.row_dimensions[row].height = 18
        row += 1

    # ── Rodapé com contagem ───────────────────────────────────
    row += 1
    total = sum(1 for p in pastas if p['nivel'] > 0)
    ws.cell(row=row, column=1,
            value=f"Total de subpastas mapeadas: {total}").font = Font(
                bold=True, italic=True, color="555555")

    wb.save(saida)
    print(f"  ✔ Excel gerado: {saida}")


# ── Main ──────────────────────────────────────────────────────
if __name__ == "__main__":
    print("\n" + "═" * 50)
    print("   Mapeador de Hierarquia de Pastas")
    print("═" * 50)

    raiz_path = pathlib.Path(PASTA_RAIZ)
    print(f"\n  Lendo: {raiz_path}")

    pastas = coletar_pastas(PASTA_RAIZ)

    print(f"  Pastas encontradas: {len(pastas)}\n")

    gerar_txt(pastas, SAIDA_TXT, raiz_path.name)
    gerar_excel(pastas, SAIDA_EXCEL, raiz_path.name)

    print(f"\n  Pronto! Arquivos salvos em:\n  {SCRIPT_DIR}\n")
